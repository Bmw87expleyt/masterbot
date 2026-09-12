import asyncio
import json
import os
import re
from datetime import datetime
from typing import Optional

import aiosqlite
import openpyxl
import speech_recognition as sr
from aiogram import BaseMiddleware, Bot, Dispatcher, F, types
from aiogram.filters import Command, CommandStart
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.storage.memory import MemoryStorage
from aiogram.types import (
    BotCommand,
    FSInputFile,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    KeyboardButton,
    ReplyKeyboardMarkup,
    TelegramObject,
)
from pydub import AudioSegment

# ======================
# CONFIG & INIT
# ======================
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PRICES_FILE = os.path.join(BASE_DIR, "prices.json")
SETTINGS_FILE = os.path.join(BASE_DIR, "settings.json")
DB_FILE = os.path.join(BASE_DIR, "users.db")
TEMP_DIR = os.path.join(BASE_DIR, "temp")
os.makedirs(TEMP_DIR, exist_ok=True)

API_TOKEN = "8704247957:AAH53nb6eBhUsnbB3YH-_AsWL4_5u8ah66E"
ADMIN_IDS = {292170708, 885001694}

bot = Bot(token=API_TOKEN)
dp = Dispatcher(storage=MemoryStorage())


# ======================
# DATABASE (aiosqlite)
# ======================
async def init_db():
    async with aiosqlite.connect(DB_FILE) as db:
        await db.execute(
            """
            CREATE TABLE IF NOT EXISTS users (
                user_id INTEGER PRIMARY KEY,
                username TEXT,
                full_name TEXT,
                phone TEXT,
                is_blocked INTEGER DEFAULT 0,
                joined_at TIMESTAMP
            )
            """
        )
        await db.commit()


async def get_user(user_id: int):
    async with aiosqlite.connect(DB_FILE) as db:
        async with db.execute(
            "SELECT * FROM users WHERE user_id = ?", (user_id,)
        ) as cur:
            return await cur.fetchone()


async def add_user(user_id: int, username: str, full_name: str, phone: str):
    async with aiosqlite.connect(DB_FILE) as db:
        await db.execute(
            """
            INSERT OR REPLACE INTO users (user_id, username, full_name, phone, is_blocked, joined_at)
            VALUES (?, ?, ?, ?, COALESCE((SELECT is_blocked FROM users WHERE user_id = ?), 0), ?)
            """,
            (user_id, username, full_name, phone, user_id, datetime.now()),
        )
        await db.commit()


async def update_block_status(user_id: int, status: int):
    async with aiosqlite.connect(DB_FILE) as db:
        await db.execute(
            "UPDATE users SET is_blocked = ? WHERE user_id = ?",
            (status, user_id),
        )
        await db.commit()


async def get_all_users():
    async with aiosqlite.connect(DB_FILE) as db:
        async with db.execute(
            "SELECT user_id, full_name, phone, is_blocked, joined_at FROM users"
        ) as cur:
            return await cur.fetchall()


# ======================
# STORAGE & PARSERS
# ======================
def load_json(path: str, default):
    if not os.path.exists(path):
        return default
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def save_json(path: str, data):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def get_settings():
    return load_json(
        SETTINGS_FILE,
        {
            "contact_text": "🛠 Специализированная лаборатория по компонентному ремонту и переклейке.\nПринимаем устройства лично и доставкой со всей РФ.",
            "admin_chat_id": None,
        },
    )


def get_prices():
    data = load_json(PRICES_FILE, {"brands": {}})
    if "brands" not in data:
        data["brands"] = {}
    return data


def add_price_item(brand: str, model: str, fault: str, price: str):
    data = get_prices()
    brand_cap = brand.strip().capitalize()
    model_cap = model.strip()

    if brand_cap not in data["brands"]:
        data["brands"][brand_cap] = {}
    if model_cap not in data["brands"][brand_cap]:
        data["brands"][brand_cap][model_cap] = {}

    data["brands"][brand_cap][model_cap][fault.strip()] = price.strip()
    save_json(PRICES_FILE, data)


def parse_service_query(raw_text: str) -> Optional[tuple]:
    text = raw_text.replace(",", " ").strip()

    if "|" in text:
        parts = [p.strip() for p in text.split("|")]
        if len(parts) >= 4:
            return parts[0], parts[1], parts[2], parts[3]

    price_match = re.search(r"(\d+)(?:\s*руб|\s*₽)?$", text)
    if not price_match:
        return None

    price = price_match.group(1)
    body = text[: price_match.start()].strip()
    words = body.split()
    if len(words) < 3:
        return None

    brand = words[0]
    fault = words[-1]
    model = " ".join(words[1:-1])
    return brand, model, fault, price


def transcribe_voice_to_text(ogg_path: str) -> str:
    wav_path = ogg_path.replace(".ogg", ".wav")
    try:
        sound = AudioSegment.from_file(ogg_path, format="ogg")
        sound.export(wav_path, format="wav")

        recognizer = sr.Recognizer()
        with sr.AudioFile(wav_path) as src:
            audio_data = recognizer.record(src)
            return recognizer.recognize_google(audio_data, language="ru-RU")
    except Exception:
        return ""
    finally:
        if os.path.exists(wav_path):
            os.remove(wav_path)


# ======================
# EXCEL ENGINE
# ======================
def generate_excel_sync(data: dict, filepath: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Прайс"
    ws.append(["Бренд", "Модель", "Услуга / Поломка", "Цена (руб)"])

    for brand, models in data.get("brands", {}).items():
        for model, faults in models.items():
            for fault, price in faults.items():
                ws.append([brand, model, fault, str(price)])

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 3
    wb.save(filepath)


def parse_excel_sync(filepath: str) -> dict:
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active
    res = {"brands": {}}

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0] or not row[1]:
            continue
        brand, model, fault, price = [
            str(x).strip() if x is not None else "" for x in row[:4]
        ]
        if brand not in res["brands"]:
            res["brands"][brand] = {}
        if model not in res["brands"][brand]:
            res["brands"][brand][model] = {}
        if fault:
            res["brands"][brand][model][fault] = price
    return res


# ======================
# KEYBOARDS & FSM
# ======================
main_kb = ReplyKeyboardMarkup(
    keyboard=[
        [KeyboardButton(text="🔍 Найти модель / Прайс")],
        [
            KeyboardButton(text="📍 Контакты и доставка"),
            KeyboardButton(text="💬 Задать вопрос мастеру"),
        ],
    ],
    resize_keyboard=True,
)

cancel_kb = ReplyKeyboardMarkup(
    keyboard=[[KeyboardButton(text="Отмена")]], resize_keyboard=True
)

phone_request_kb = ReplyKeyboardMarkup(
    keyboard=[[KeyboardButton(text="📞 Отправить телефон", request_contact=True)]],
    resize_keyboard=True,
)

admin_kb = ReplyKeyboardMarkup(
    keyboard=[
        [
            KeyboardButton(text="🎙 Надиктовать позицию (Голос/Текст)"),
            KeyboardButton(text="📥 Скачать прайс (Excel)"),
        ],
        [
            KeyboardButton(text="👥 База клиентов"),
            KeyboardButton(text="🚫 Блокировка"),
        ],
        [KeyboardButton(text="⬅️ Выйти из админки")],
    ],
    resize_keyboard=True,
)


class RegFlow(StatesGroup):
    name = State()
    phone = State()


class SearchFlow(StatesGroup):
    waiting_query = State()


class QuestionFlow(StatesGroup):
    waiting_text = State()


class AdminFlow(StatesGroup):
    waiting_add_item = State()
    waiting_block_id = State()
    waiting_reply_text = State()


# ======================
# MIDDLEWARE
# ======================
class AuthMiddleware(BaseMiddleware):
    async def __call__(self, handler, event: TelegramObject, data: dict):
        if not isinstance(event, types.Message) or not event.from_user:
            return await handler(event, data)

        state: FSMContext = data.get("state")
        current_state = await state.get_state() if state else None

        if current_state in [RegFlow.name.state, RegFlow.phone.state]:
            return await handler(event, data)

        user = await get_user(event.from_user.id)
        if not user:
            if state:
                await state.set_state(RegFlow.name)
            await event.answer(
                "Здравствуйте! 👋\nДля доступа к боту укажите ваше <b>Имя</b>:",
                parse_mode="HTML",
            )
            return

        if user[4] == 1:
            await event.answer("⛔️ Доступ ограничен.")
            return

        data["db_user"] = user
        return await handler(event, data)


dp.message.outer_middleware(AuthMiddleware())


# ======================
# CLIENT REGISTRATION
# ======================
@dp.message(RegFlow.name)
async def reg_name(message: types.Message, state: FSMContext):
    if not message.text:
        return await message.answer("Введите имя текстовым сообщением.")
    await state.update_data(full_name=message.text.strip())
    await state.set_state(RegFlow.phone)
    await message.answer(
        f"Принято, {message.text}!\nНажмите кнопку ниже для подтверждения телефона:",
        reply_markup=phone_request_kb,
    )


@dp.message(RegFlow.phone)
async def reg_phone(message: types.Message, state: FSMContext):
    if not message.contact:
        return await message.answer(
            "Воспользуйтесь кнопкой «📞 Отправить телефон»."
        )

    data = await state.get_data()
    u = message.from_user
    await add_user(
        u.id, u.username or "", data["full_name"], message.contact.phone_number
    )
    await state.clear()
    await message.answer(
        "✅ Регистрация пройдена.", reply_markup=main_kb
    )

    alert = (
        f"👤 <b>Новый пользователь:</b>\n"
        f"Имя: {data['full_name']}\n"
        f"Тел: {message.contact.phone_number}\n"
        f"ID: <code>{u.id}</code>\n"
        f"Username: @{u.username or 'отсутствует'}"
    )
    for adm in ADMIN_IDS:
        try:
            await bot.send_message(adm, alert, parse_mode="HTML")
        except Exception:
            pass


# ======================
# CLIENT LOGIC
# ======================
@dp.message(CommandStart())
async def start(message: types.Message, state: FSMContext):
    await state.clear()
    await message.answer("Главное меню:", reply_markup=main_kb)


@dp.message(F.text == "📍 Контакты и доставка")
async def contacts(message: types.Message):
    cfg = get_settings()
    await message.answer(
        f"📍 <b>Информация:</b>\n\n{cfg.get('contact_text')}",
        reply_markup=main_kb,
        parse_mode="HTML",
    )


@dp.message(F.text.in_(["Отмена", "⬅️ Выйти из админки"]))
async def cancel_handler(message: types.Message, state: FSMContext):
    await state.clear()
    await message.answer("Отменено.", reply_markup=main_kb)


@dp.message(F.text == "💬 Задать вопрос мастеру")
@dp.message(Command("ask"))
async def ask_start(message: types.Message, state: FSMContext):
    await state.set_state(QuestionFlow.waiting_text)
    await message.answer(
        "✍️ Опишите модель и суть проблемы (можно отправить голосовое):",
        reply_markup=cancel_kb,
    )


@dp.message(QuestionFlow.waiting_text, F.voice)
@dp.message(QuestionFlow.waiting_text, F.text)
async def ask_process(message: types.Message, state: FSMContext, db_user: tuple):
    query_text = message.text
    if message.voice:
        file_id = message.voice.file_id
        file = await bot.get_file(file_id)
        voice_path = os.path.join(TEMP_DIR, f"{file_id}.ogg")
        await bot.download_file(file.file_path, voice_path)
        query_text = await asyncio.to_thread(
            transcribe_voice_to_text, voice_path
        )
        if os.path.exists(voice_path):
            os.remove(voice_path)

    admin_kb_reply = InlineKeyboardMarkup(
        inline_keyboard=[
            [
                InlineKeyboardButton(
                    text="↩️ Ответить клиенту",
                    callback_data=f"reply_to_{db_user[0]}",
                )
            ]
        ]
    )

    alert = (
        f"🚨 <b>Новый запрос мастеру!</b>\n\n"
        f"От: {db_user[2]} (<code>{db_user[0]}</code>)\n"
        f"Тел: {db_user[3]}\n"
        f"Сообщение: {query_text or 'Аудиосообщение (не распознано)'}"
    )

    for adm in ADMIN_IDS:
        try:
            await bot.send_message(
                adm, alert, reply_markup=admin_kb_reply, parse_mode="HTML"
            )
        except Exception:
            pass

    await state.clear()
    await message.answer(
        "✅ Вопрос передан мастеру. Ответ поступит в этот диалог.",
        reply_markup=main_kb,
    )


@dp.message(F.text == "🔍 Найти модель / Прайс")
async def search_init(message: types.Message, state: FSMContext):
    await state.set_state(SearchFlow.waiting_query)
    await message.answer(
        "Введите модель (например, <code>iPhone 13</code> или <code>S22 Ultra</code>):",
        reply_markup=cancel_kb,
        parse_mode="HTML",
    )


@dp.message(SearchFlow.waiting_query)
async def search_process(message: types.Message, state: FSMContext):
    query = message.text.strip().lower()
    data = await asyncio.to_thread(get_prices)
    results = []

    for brand, models in data.get("brands", {}).items():
        for model, faults in models.items():
            full_title = f"{brand} {model}".lower()
            if query in full_title or any(
                q in full_title for q in query.split()
            ):
                results.append((brand, model, faults))

    await state.clear()
    if not results:
        return await message.answer(
            f"По запросу «{message.text}» ничего не найдено.\nИспользуйте кнопку «💬 Задать вопрос мастеру».",
            reply_markup=main_kb,
        )

    for brand, model, faults in results[:5]:
        card = f"📱 <b>{brand} {model}</b>\n\n"
        for fault, pr in faults.items():
            card += f"▪️ {fault}: <b>{pr} ₽</b>\n"
        await message.answer(card, parse_mode="HTML")

    await message.answer("Готово.", reply_markup=main_kb)


# ======================
# ADMIN: VOICE ENTRY & REPLIES
# ======================
@dp.message(Command("admin"), F.from_user.id.in_(ADMIN_IDS))
async def admin_panel(message: types.Message, state: FSMContext):
    await state.clear()
    await message.answer("Панель администратора:", reply_markup=admin_kb)


@dp.message(
    F.text == "🎙 Надиктовать позицию (Голос/Текст)",
    F.from_user.id.in_(ADMIN_IDS),
)
async def admin_voice_add_start(message: types.Message, state: FSMContext):
    await state.set_state(AdminFlow.waiting_add_item)
    await message.answer(
        "Надиктуйте голосовое или напишите текстом в формате:\n"
        "<code>Бренд Модель Услуга Цена</code>\n\n"
        "Например:\n"
        "<i>«Samsung S22 Ultra переклейка 6000»</i>",
        reply_markup=cancel_kb,
        parse_mode="HTML",
    )


@dp.message(AdminFlow.waiting_add_item, F.voice, F.from_user.id.in_(ADMIN_IDS))
@dp.message(AdminFlow.waiting_add_item, F.text, F.from_user.id.in_(ADMIN_IDS))
async def admin_voice_add_process(message: types.Message, state: FSMContext):
    raw_text = message.text

    if message.voice:
        status_msg = await message.answer("⏳ Обрабатываю голос...")
        file_id = message.voice.file_id
        file = await bot.get_file(file_id)
        voice_path = os.path.join(TEMP_DIR, f"{file_id}.ogg")
        await bot.download_file(file.file_path, voice_path)

        raw_text = await asyncio.to_thread(
            transcribe_voice_to_text, voice_path
        )
        if os.path.exists(voice_path):
            os.remove(voice_path)
        await status_msg.delete()

    if not raw_text:
        return await message.answer(
            "Не удалось разобрать аудио. Надиктуйте четче или пришлите текстом."
        )

    parsed = parse_service_query(raw_text)
    if not parsed:
        return await message.answer(
            f"Распознано: «<i>{raw_text}</i>»\n\n"
            "Не удалось выделить параметры. Формат: <code>Бренд Модель Ремонт Цена</code>.",
            parse_mode="HTML",
        )

    brand, model, fault, price = parsed
    await asyncio.to_thread(add_price_item, brand, model, fault, price)
    await state.clear()
    await message.answer(
        f"✅ <b>Добавлено в базу:</b>\n\n"
        f"Бренд: <b>{brand}</b>\n"
        f"Модель: <b>{model}</b>\n"
        f"Услуга: <b>{fault}</b>\n"
        f"Цена: <b>{price} ₽</b>",
        reply_markup=admin_kb,
        parse_mode="HTML",
    )


@dp.callback_query(F.data.startswith("reply_to_"), F.from_user.id.in_(ADMIN_IDS))
async def admin_reply_start(cb: types.CallbackQuery, state: FSMContext):
    client_id = int(cb.data.split("_")[2])
    await state.set_state(AdminFlow.waiting_reply_text)
    await state.update_data(target_client=client_id)
    await cb.message.answer(
        f"Ответ клиенту <code>{client_id}</code> (текстом или голосом):",
        reply_markup=cancel_kb,
        parse_mode="HTML",
    )
    await cb.answer()


@dp.message(AdminFlow.waiting_reply_text, F.from_user.id.in_(ADMIN_IDS))
async def admin_reply_send(message: types.Message, state: FSMContext):
    data = await state.get_data()
    client_id = data.get("target_client")
    reply_content = message.text

    if message.voice:
        file_id = message.voice.file_id
        file = await bot.get_file(file_id)
        voice_path = os.path.join(TEMP_DIR, f"{file_id}.ogg")
        await bot.download_file(file.file_path, voice_path)
        reply_content = await asyncio.to_thread(
            transcribe_voice_to_text, voice_path
        )
        if os.path.exists(voice_path):
            os.remove(voice_path)

    try:
        await bot.send_message(
            client_id,
            f"👨‍🔧 <b>Ответ мастера:</b>\n\n{reply_content}",
            parse_mode="HTML",
        )
        await message.answer("✅ Ответ передан клиенту.", reply_markup=admin_kb)
    except Exception as e:
        await message.answer(f"❌ Ошибка отправки: {e}", reply_markup=admin_kb)

    await state.clear()


# ======================
# EXCEL & MANAGEMENT
# ======================
@dp.message(F.text == "📥 Скачать прайс (Excel)", F.from_user.id.in_(ADMIN_IDS))
async def admin_export_excel(message: types.Message):
    filepath = os.path.join(BASE_DIR, "prices.xlsx")
    data = await asyncio.to_thread(get_prices)
    await asyncio.to_thread(generate_excel_sync, data, filepath)
    await message.answer_document(
        FSInputFile(filepath), caption="Текущий файл прайс-листа."
    )


@dp.message(F.document, F.from_user.id.in_(ADMIN_IDS))
async def admin_import_excel(message: types.Message):
    if not message.document.file_name.endswith(".xlsx"):
        return await message.answer("Файл должен иметь расширение .xlsx")

    dest = os.path.join(TEMP_DIR, "import_prices.xlsx")
    await bot.download(message.document, destination=dest)

    try:
        new_data = await asyncio.to_thread(parse_excel_sync, dest)
        save_json(PRICES_FILE, new_data)
        await message.answer("✅ Прайс успешно импортирован.")
    except Exception as e:
        await message.answer(f"❌ Ошибка парсинга таблицы: {e}")
    finally:
        if os.path.exists(dest):
            os.remove(dest)


@dp.message(F.text == "👥 База клиентов", F.from_user.id.in_(ADMIN_IDS))
async def admin_list_users(message: types.Message):
    users = await get_all_users()
    if not users:
        return await message.answer("База клиентов пуста.")

    text = "👥 <b>Клиенты:</b>\n\n"
    for u in users:
        st = "🔴 Блок" if u[4] else "🟢 Доступ"
        text += f"ID: <code>{u[0]}</code> | {u[2]} | Тел: {u[3]} | {st}\n"

    for chunk in [text[i : i + 3900] for i in range(0, len(text), 3900)]:
        await message.answer(chunk, parse_mode="HTML")


@dp.message(F.text == "🚫 Блокировка", F.from_user.id.in_(ADMIN_IDS))
async def admin_ban_menu(message: types.Message):
    kb = InlineKeyboardMarkup(
        inline_keyboard=[
            [
                InlineKeyboardButton(
                    text="Заблокировать", callback_data="adm_ban"
                ),
                InlineKeyboardButton(
                    text="Разблокировать", callback_data="adm_unban"
                ),
            ]
        ]
    )
    await message.answer("Действие со статусом клиента:", reply_markup=kb)


@dp.callback_query(
    F.data.in_(["adm_ban", "adm_unban"]), F.from_user.id.in_(ADMIN_IDS)
)
async def admin_ban_select(cb: types.CallbackQuery, state: FSMContext):
    action = 1 if cb.data == "adm_ban" else 0
    await state.update_data(ban_action=action)
    await state.set_state(AdminFlow.waiting_block_id)
    await cb.message.answer(
        "Укажите числовой ID пользователя:", reply_markup=cancel_kb
    )
    await cb.answer()


@dp.message(AdminFlow.waiting_block_id, F.from_user.id.in_(ADMIN_IDS))
async def admin_ban_process(message: types.Message, state: FSMContext):
    if not message.text.isdigit():
        return await message.answer("ID должен состоять только из цифр.")
    data = await state.get_data()
    uid = int(message.text)
    await update_block_status(uid, data["ban_action"])
    await state.clear()
    state_str = "заблокирован" if data["ban_action"] == 1 else "разблокирован"
    await message.answer(
        f"Пользователь <code>{uid}</code> {state_str}.",
        reply_markup=admin_kb,
        parse_mode="HTML",
    )


# ======================
# RUNNER
# ======================
async def on_startup():
    await init_db()
    await bot.set_my_commands(
        [
            BotCommand(command="start", description="Главное меню"),
            BotCommand(command="ask", description="Задать вопрос мастеру"),
            BotCommand(command="admin", description="Панель управления"),
        ]
    )


async def main():
    if not os.path.exists(SETTINGS_FILE):
        save_json(SETTINGS_FILE, get_settings())
    if not os.path.exists(PRICES_FILE):
        save_json(PRICES_FILE, {"brands": {}})

    dp.startup.register(on_startup)
    await dp.start_polling(bot)


if __name__ == "__main__":
    asyncio.run(main())
